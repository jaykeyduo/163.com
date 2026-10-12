#!/usr/bin/env python3
"""金山 2027 物流成本预算框架（驱动表，不填假精确金额）。"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pathlib import Path

OUT = Path("/workspace/deliverables/金山_2027物流成本预算框架.xlsx")

THIN = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)
FILL_H = PatternFill("solid", fgColor="1F4E79")
FILL_S = PatternFill("solid", fgColor="D6EAF8")
FILL_Y = PatternFill("solid", fgColor="FFF4CC")
FILL_G = PatternFill("solid", fgColor="E8F5E9")
FILL_R = PatternFill("solid", fgColor="FDECEA")
FONT_H = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
FONT_B = Font(bold=True, name="Calibri", size=11)
FONT_N = Font(name="Calibri", size=11)
WRAP = Alignment(wrap_text=True, vertical="center")


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = FILL_H
        cell.font = FONT_H
        cell.alignment = WRAP
        cell.border = THIN


def style_body(ws, r1, r2, cols):
    for r in range(r1, r2 + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(r, c)
            cell.font = FONT_N
            cell.alignment = WRAP
            cell.border = THIN


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_note(ws, row, cols, text, fill=FILL_Y):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, text)
    cell.fill = fill
    cell.font = FONT_N
    cell.alignment = WRAP
    for c in range(1, cols + 1):
        ws.cell(row, c).border = THIN
        ws.cell(row, c).fill = fill


wb = Workbook()

# --- 0 读我 ---
ws = wb.active
ws.title = "00-读我"
ws["A1"] = "金山工厂物流成本预算框架（给财务四类：运费 / 外仓 / 进口手续费 / 其他）"
ws["A1"].font = Font(bold=True, size=14, name="Calibri", color="1F4E79")
ws.merge_cells("A1:B1")
rows = [
    ("用途", "正确往下做月度预算的驱动结构。产量未按月锁定前，不填假精确 RMB。"),
    ("财务假设", "2026 年底现有产线或迁独山；2027 开启 Chromate & Silicone 本地生产。"),
    ("产量口径（已有、待 MP 月分解）", "Alex 2026-08：含硅 72 T/年 + 含铬 400 T/年 = 472 T/年成品，约 39 T/月（均匀假设）。"),
    ("真正卡住运费的缺口", "472 T 是成品产量，不是进口原料吨。缺：BOM、进口 vs 国产原料占比、到港卸点（洋山→金山 还是 先独山再调拨）。"),
    ("CIF 口径", "德→中空运/海运由德国指定、CIF。国际段运费通常在采购价里，不进中国「运费」科目。中国运费=到港后内陆+国内配送+独山↔金山调拨。须财务书面确认。"),
    ("外仓口径（已定）", "金山自有仓，不租外仓。本列填 0。叉车采购与仓操作人员进「其他」，不要塞进外仓。长春/沧州外仓仅汽车，C&S 成品是否走外仓待销售确认，默认不进本预算。"),
    ("两套时间", "2026 Q4：现有产线迁独山=一次性搬迁物流，不是 2027 运行费率。2027：C&S 投产月度运行成本。"),
    ("事实源", "projects/2026-08-jinshan-wh-restart.md；knowledge/01-company、03-network、04-operations。"),
]
ws["A2"] = "项"
ws["B2"] = "说明"
style_header(ws, 2, 2)
for i, (a, b) in enumerate(rows, 3):
    ws.cell(i, 1, a)
    ws.cell(i, 2, b)
style_body(ws, 3, 2 + len(rows), 2)
ws.row_dimensions[1].height = 24
for r in range(3, 11):
    ws.row_dimensions[r].height = 48
set_widths(ws, [28, 92])

# --- 1 驱动 ---
ws = wb.create_sheet("01-产量与驱动")
headers = ["驱动", "建议值 / 公式", "状态", "Owner", "不填会怎样"]
ws.append(headers)
style_header(ws, 1, 5)
drivers = [
    ["成品年产量 含硅 T", "72", "已有（Alex 2026-08）", "生产", "无月计划则只能均匀摊 6 T/月"],
    ["成品年产量 含铬 T", "400", "已有（Alex 2026-08）", "生产", "均匀摊 ~33 T/月"],
    ["合计成品 T/年", "=B2+B3", "由上两项", "生产", ""],
    ["2027 月产量曲线", "待填 12 个数；暂可 =B4/12", "缺", "计划 MP", "运费/手续费无法按月"],
    ["原料投入系数（原料T / 成品T）", "待工艺 BOM", "缺", "生产/工艺", "有成品吨仍算不出原料吨"],
    ["进口原料占比 %", "待采购：德国进口 vs 国内采购", "缺", "采购", "进口运费与手续费无法算"],
    ["进口卸点", "洋山直送金山 / 先独山再调拨（二选一）", "缺", "物流+MP", "内陆运费差一倍搬运"],
    ["CIF 国际运费是否进中国运费科目", "建议否（含在采购价）", "待财务书面", "财务", "与德国重复计算或漏计"],
    ["关税/增值税是否进本表", "建议否（单独税项，不进四类）", "待财务书面", "财务", "四类被税撑爆，无法比物流"],
    ["成品是否经长春/沧州外仓", "建议否（C&S 非汽车卫星仓默认）", "待销售", "销售", "外仓列可能不该为 0"],
    ["金山→客户配送模式", "公路危货零担（国内主模式）", "网络已运行", "物流", "缺客户结构则吨公里未知"],
    ["仓常驻人数", "2（方案 A）；或 3（方案 B 不依赖独山 Backup）", "物流已建议", "物流/厂长", "其他-人力无法编"],
    ["独山 Backup 人天/月", "待厂长锁定上限", "缺", "厂长", "成本中心扯皮"],
    ["新前移式 1.4T", "1 台新购", "物流已建议", "物流/工程", "CapEx vs 折旧问财务"],
    ["新平衡重", "1 台新购；吨位待工艺", "物流已建议", "生产定吨位", "错买/晚买"],
    ["2015 前移", "过渡备用，不替代新车", "现役", "物流", "勿把旧车当产能"],
]
for row in drivers:
    ws.append(row)
style_body(ws, 2, 1 + len(drivers), 5)
ws["B4"] = "=B2+B3"
ws["C4"].fill = FILL_G
add_note(
    ws,
    18,
    5,
    "公式（有数后再填金额）：月进口原料吨 = 月成品吨 × 原料投入系数 × 进口原料占比。中国侧进口运费驱动=该吨数（到港后内陆）。手续费驱动=进口票数或票+吨（代理报价）。国内配送驱动=月成品吨×平均运距。",
)
ws.row_dimensions[18].height = 48
set_widths(ws, [32, 42, 22, 16, 28])
ws.freeze_panes = "A2"

# --- 2 映射 ---
ws = wb.create_sheet("02-四类科目映射")
ws.append(["财务四类", "子项", "进本列？", "CapEx/OpEx", "物流建议", "缺什么才能填 RMB"])
style_header(ws, 1, 6)
maps = [
    ["运费", "CIF 汉堡/FRA→洋山/PVG 国际段", "默认否", "在采购价", "德国指定承运；中国物流不订舱。财务若要集团全成本，请另开「CIF 国际」备查列，不要与中国运费混加。", "财务确认科目"],
    ["运费", "洋山/PVG→金山 进口内陆（危货公路）", "是", "OpEx", "C&S 投产后的主进口运费。若先入独山再调拨，改记调拨运费。", "月进口吨 + 卸点 + 危货费率"],
    ["运费", "金山→中国客户 成品配送", "是", "OpEx", "国内危货零担；温控如涉及 UV/水性另列。", "客户区域结构 + 月成品吨"],
    ["运费", "金山↔独山 原料/半成品/包材调拨", "是（若发生）", "OpEx", "共用独山库存或包材时必有。含铬分区后尽量少调。", "调拨规则 + 月吨"],
    ["运费", "2026 现有产线迁独山（设备/剩余库存）", "一次性，可单列 2026-12", "一次性", "不是 2027 运行费率。危货搬迁+特种设备。", "搬迁清单吨位/车次"],
    ["运费", "紧急空运进口", "是（例外）", "OpEx", "不编进基线；单列应急。空运亦 CIF、德国指定。", "政策：是否允许"],
    ["外仓", "金山厂内仓租赁", "否=0", "—", "自有空间，不租。", "无"],
    ["外仓", "长春/沧州 3PL", "默认否=0", "OpEx 若启用", "现网仅汽车。C&S 默认不进卫星仓。", "销售是否要求外仓"],
    ["外仓", "港区/危库暂存超期", "其他或运费附带", "OpEx 异常", "设计为零；超期再进异常。", "不编基线"],
    ["进口手续费", "报关代理 / 换单 / 港杂（CIF 到港后买方段）", "是", "OpEx", "CIF 不包含中国报关与港后费用。", "票数或吨 + 代理费率"],
    ["进口手续费", "危化进口申报 / 查验陪同", "是", "OpEx", "含铬/含硅查验概率高于普通漆，费率可高于独山常规。", "代理报价"],
    ["进口手续费", "关税、增值税、消费税", "建议不进本四类", "税", "请财务进税项；物流表只留手续费。", "HS + 完税价格"],
    ["其他", "仓操作人员工资社保（常驻 2 人）", "是", "OpEx", "方案 A。方案 B=3 人则改 3。", "HR 人薪"],
    ["其他", "独山 Backup 人天（费用进金山成本中心）", "是", "OpEx", "必须制度化，避免独山加班白送。", "人天上限 × 日费率"],
    ["其他", "新前移 1.4T + 新平衡重", "是（购入月现金或月折旧，二选一问财务）", "CapEx", "2012 旧车不计产能；2015 仅备用。", "询价 + 吨位 + 折旧年限"],
    ["其他", "叉车维保 / 充电 / 备件", "是", "OpEx", "工程主责，预算仍要列。", "维保合同"],
    ["其他", "10 号一楼办公（电脑+打印）", "是", "小 CapEx/OpEx", "最小闭环。", "询价"],
    ["其他", "保洁（仓+办公）", "是（若物流承担）", "OpEx", "主责工厂公共则本列 0，须书面。", "主责部门"],
    ["其他", "含铬/含硅分区、标识、PPE、持证培训", "是或 EHS 科目", "CapEx/OpEx", "合规门禁；无路径则推迟投产。不填 0 假装没有。", "EHS 清单"],
    ["其他", "废弃物（含铬相关）", "仅当仓/物流主责", "OpEx", "生产提出过；默认生产/EHS 主责，仓不兜底。", "厂长点名主责"],
    ["其他", "出单打签系统/耗材", "否（不默认进仓）", "OpEx 在 Shipping/计划", "仓员不顺便做。", "主责部门"],
]
for row in maps:
    ws.append(row)
    last = ws.max_row
    if row[2].startswith("否") or "默认否" in row[2]:
        ws.cell(last, 3).fill = FILL_G
    if "缺" in str(row[5]) or row[5] not in ("无",):
        if row[5] != "无":
            ws.cell(last, 6).fill = FILL_Y
style_body(ws, 2, 1 + len(maps), 6)
set_widths(ws, [14, 42, 22, 14, 48, 28])
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:F{1 + len(maps)}"

# --- 3 2027 月度 ---
ws = wb.create_sheet("03-2027月度四类")
months = [f"2027-{m:02d}" for m in range(1, 13)]
ws.append(["科目", "子项 / 驱动"] + months + ["2027合计", "填数规则"])
style_header(ws, 1, 16)
lines = [
    ("驱动", "成品产量 T（默认 472/12，MP 可改）", "flat"),
    ("驱动", "进口原料 T（=产量×投入系数×进口占比；缺系数则留空）", "empty"),
    ("运费", "到港后内陆 洋山/PVG→金山", "empty"),
    ("运费", "金山→客户 成品配送", "empty"),
    ("运费", "金山↔独山 调拨", "empty"),
    ("运费", "紧急空运（基线 0）", "zero"),
    ("运费", "运费小计", "sum_freight"),
    ("外仓", "金山厂内租赁（自有=0）", "zero"),
    ("外仓", "长春/沧州 3PL（默认 0）", "zero"),
    ("外仓", "外仓小计", "sum_wh"),
    ("进口手续费", "报关/换单/港杂", "empty"),
    ("进口手续费", "危化申报/查验", "empty"),
    ("进口手续费", "进口手续费小计", "sum_imp"),
    ("其他", "仓操作 2 人（月薪资，HR 填）", "empty"),
    ("其他", "独山 Backup 人天", "empty"),
    ("其他", "叉车折旧或购入月现金（与财务选口径）", "empty"),
    ("其他", "叉车维保", "empty"),
    ("其他", "办公+保洁", "empty"),
    ("其他", "含铬硅合规/PPE/培训（无报价则留空）", "empty"),
    ("其他", "其他小计", "sum_oth"),
    ("合计", "四类合计", "sum_all"),
]
for i, (cat, name, kind) in enumerate(lines):
    r = i + 2
    ws.cell(r, 1, cat)
    ws.cell(r, 2, name)
    if kind == "flat":
        for c in range(3, 15):
            ws.cell(r, c, "=驱动!产量")  # placeholder overwritten below
            ws.cell(r, c, f"='01-产量与驱动'!$B$4/12")
        ws.cell(r, 15, f"=SUM(C{r}:N{r})")
        ws.cell(r, 16, "产量驱动行。MP 给爬坡后改各月，不要改合计公式。")
        for c in range(3, 16):
            ws.cell(r, c).fill = FILL_S
            ws.cell(r, c).number_format = "0.00"
    elif kind == "zero":
        for c in range(3, 15):
            ws.cell(r, c, 0)
        ws.cell(r, 15, f"=SUM(C{r}:N{r})")
        ws.cell(r, 16, "基线 0。启用外仓/空运时改对应月。")
        for c in range(3, 16):
            ws.cell(r, c).fill = FILL_G
            ws.cell(r, c).number_format = "#,##0"
    elif kind == "empty":
        for c in range(3, 15):
            ws.cell(r, c, None)
        ws.cell(r, 15, f"=SUM(C{r}:N{r})")
        ws.cell(r, 16, "有费率×驱动后再填。禁止估一个整数充数。")
        for c in range(3, 16):
            ws.cell(r, c).fill = FILL_Y
            ws.cell(r, c).number_format = "#,##0"
    elif kind.startswith("sum"):
        # row index mapping
        pass
    ws.row_dimensions[r].height = 32

# 小计公式：行号 运费=8, 外仓=11, 手续=14, 其他=21, 合计=22 （header=1, 驱动2-3, 运费4-8, 外仓9-11, 手续12-14, 其他15-21, 合计22）
# lines: 1驱动产量 r2, 2进口原料 r3, 3-6运费明细 r4-7, 7运费小计 r8,
# 8-9外仓 r9-10, 10外仓小计 r11, 11-12手续 r12-13, 13手续小计 r14,
# 14-19其他 r15-20, 20其他小计 r21, 21合计 r22
for c in range(3, 15):
    cl = get_column_letter(c)
    ws.cell(8, c, f"={cl}4+{cl}5+{cl}6+{cl}7")
    ws.cell(11, c, f"={cl}9+{cl}10")
    ws.cell(14, c, f"={cl}12+{cl}13")
    ws.cell(21, c, f"={cl}15+{cl}16+{cl}17+{cl}18+{cl}19+{cl}20")
    ws.cell(22, c, f"={cl}8+{cl}11+{cl}14+{cl}21")
for r in (8, 11, 14, 21, 22):
    ws.cell(r, 15, f"=SUM(C{r}:N{r})")
    for c in range(1, 17):
        ws.cell(r, c).font = FONT_B
        ws.cell(r, c).fill = FILL_S if r != 22 else FILL_H
        if r == 22:
            ws.cell(r, c).font = FONT_H
    ws.cell(r, 16, "小计，勿手改")
    ws.cell(r, 3).number_format = "#,##0"

ws.cell(8, 16, "运费小计")
ws.cell(11, 16, "外仓小计（自有仓应为 0）")
ws.cell(14, 16, "进口手续费小计（不含关税增值税）")
ws.cell(21, 16, "其他小计")
ws.cell(22, 16, "四类合计")

style_body(ws, 2, 22, 16)
# re-apply fills after style_body
for r in (8, 11, 14, 21):
    for c in range(1, 17):
        ws.cell(r, c).fill = FILL_S
        ws.cell(r, c).font = FONT_B
for c in range(1, 17):
    ws.cell(22, c).fill = FILL_H
    ws.cell(22, c).font = FONT_H
for r in (2,):
    for c in range(3, 16):
        ws.cell(r, c).fill = FILL_S
for r in (7, 9, 10):
    for c in range(3, 15):
        ws.cell(r, c).fill = FILL_G

add_note(
    ws,
    24,
    16,
    "黄色格=待费率。绿色格=已定性为 0。蓝色小计勿手改。产量行已用 472/12；MP 给爬坡后只改第 2 行各月。进口原料行在 01 表系数齐备前保持空白。",
)
ws.row_dimensions[24].height = 36
set_widths(ws, [12, 42] + [12] * 12 + [12, 36])
ws.freeze_panes = "C2"
ws.row_dimensions[1].height = 22

# --- 4 2026 一次性 ---
ws = wb.create_sheet("04-2026迁线一次性")
ws.append(["项", "是否物流预算", "说明", "RMB（待填）", "Owner"])
style_header(ws, 1, 5)
once = [
    ["现有金山产线设备迁独山（特种/危货运输）", "是-运费一次性", "财务邮件假设 2026 年底。与 2027 C&S 运行运费分开。", None, "物流+生产"],
    ["剩余原料/成品迁独山", "是-运费一次性", "库存盘点后再估车次。", None, "仓+MP"],
    ["金山仓清场、旧叉车报废/处置", "其他一次性", "2012 两台停用车不计入 2027 产能。", None, "工程/仓"],
    ["独山接收增加的仓容/人力冲击", "不进金山四类，须书面提示财务", "独山 5 号甲类已约 80–85%。迁线可能挤独山，不是金山外仓费。", None, "厂长/财务"],
    ["2027 投产前金山仓空窗期（人已到岗、货未到）", "其他-人力", "招聘早于投产则 2026Q4–2027Q1 仍有工资。", None, "HR/物流"],
]
for row in once:
    ws.append(row)
style_body(ws, 2, 1 + len(once), 5)
ws.cell(5, 3).fill = FILL_R
set_widths(ws, [42, 28, 62, 16, 14])
ws.row_dimensions[5].height = 40
for r in range(2, 7):
    ws.row_dimensions[r].height = 36

# --- 5 提问 ---
ws = wb.create_sheet("05-待确认清单")
ws.append(["#", "问谁", "问题", "卡住哪一类", "状态"])
style_header(ws, 1, 5)
qs = [
    [1, "计划 MP / 生产", "2027 年 1–12 月含硅、含铬成品产量（可先给爬坡再给稳定月）", "运费、手续费", "未开始"],
    [2, "生产 / 工艺", "单位成品的原料投入系数；含铬/含硅分列", "运费、手续费", "未开始"],
    [3, "采购", "原料中德国进口 vs 国内采购的吨或金额占比", "运费、手续费", "未开始"],
    [4, "财务 Yiwen", "CIF 国际运费是否进中国「运费」？关税增值税是否进本四类？叉车按购入月现金还是折旧？", "全部口径", "未开始"],
    [5, "物流+MP", "进口到港后直送金山，还是先独山再调拨？", "运费", "未开始"],
    [6, "销售", "C&S 成品是否进长春/沧州外仓？主要客户区域？", "外仓、运费", "未开始"],
    [7, "厂长", "仓编制 A=2+Backup 还是 B=3？Backup 人天与成本中心？", "其他", "未开始"],
    [8, "生产", "平衡重吨位（单批投料+釜重）与投产月", "其他-叉车", "未开始"],
    [9, "EHS", "含铬存储分区、PPE、培训、废弃物主责与概算", "其他", "未开始"],
    [10, "HR", "仓员 2 人到位月、薪资包", "其他", "未开始"],
]
for row in qs:
    ws.append(row)
style_body(ws, 2, 1 + len(qs), 5)
dv = DataValidation(type="list", formula1='"未开始,进行中,已回复"', allow_blank=False)
dv.add("E2:E11")
ws.add_data_validation(dv)
set_widths(ws, [6, 18, 62, 18, 12])
for r in range(2, 12):
    ws.row_dimensions[r].height = 32

wb.save(OUT)
print(f"wrote {OUT}")
