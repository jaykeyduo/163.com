#!/usr/bin/env python3
"""Build 出口 Initiate 跨部门 Excel 勾选表 from locked v1 rules."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "deliverables" / "出口Initiate跨部门Checklist.xlsx"

# 部门, 编号, 检查项, 合格标准, 默认灯规则, 本阶段是否适用
ROWS = [
    ("业务 / PC", "B1", "目的国、最终收货人、客户还是汉堡/关联方", "国别、全称、地址类型书面", "缺国别=红", True),
    ("业务 / PC", "B2", "贸易条款（Incoterms）——中国出境报关所需", "条款书面；运保与报关一致。不卡目的国 VAT", "完全未定=红", True),
    ("业务 / PC", "B3", "10 位产品码；固化剂/稀释剂分票还是一套", "码已定义；配套品身份不混报", "无码=红", True),
    ("业务 / PC", "B4", "数量、包装意向、样品或商业、货值币种", "可报关；样品仍有 customs value", "无货值=黄（报关前绿）", True),
    ("业务 / PC", "B5", "客户交期是否把出口当国内交期承诺", "未用国内交期锁死出口票", "已锁死国内交期=红", True),
    ("业务 / PC", "B6", "S1 Export Request Sheet：产品+包装", "已交；出口桶码≠内贸", "未填=红", True),
    ("业务 / PC", "B7", "客户额外包装/标签/托盘/文件要求", "非标有主责，不默认仓配兜底", "非标无主责=黄", True),
    ("采购", "P1", "S2 是否已有该料号标准出口包装", "有 S2；新包装默认不推荐", "无 S2 且要新包装=红", True),
    ("采购", "P2", "S3 性能单抬头是否为独山", "与鉴定/危包证一致；禁止金山抬头出口", "非独山=红", True),
    ("采购", "P3", "海/空/海+空；性能单是否覆盖该模式", "采购书面适用性声明；海运证不可套空运", "模式未声明=红", True),
    ("采购", "P4", "铁桶+外箱齐套、UN 与证一致、有效期内", "缺纸箱证不得称齐套", "缺证=红；有补购 ETA=黄", True),
    ("采购", "P5", "库存是否够本票；补购 ETA", "数量+ETA；到货带 UN+年份", "无 ETA=红", True),
    ("采购", "P6", "是否 UN 年切换（大口隔离/小口报废）", "已通知仓/生产/计划/出口", "年切换未通知=黄", True),
    ("Order Clearing", "L1", "SDS 第14章 + 中文件/中文译件", "受控版本与料号一致（129号报检）", "无 SDS=红", True),
    ("Order Clearing", "L2", "危险公示标签样本（中国出口报检）", "与运输菱形标分层、不混贴", "无样本=黄（报检前绿）", True),
    ("Order Clearing", "L3", "运输标签只按鉴定+SDS14", "不因 GHS05 自动加 8 类副标", "与鉴定冲突=红", True),
    ("EHS", "E1", "是否《危险化学品目录》内（如 2828）", "书面是/否；是则走 129 号产地检验", "未判定=红", True),
    ("EHS", "E2", "分类鉴定有效；UN/类/包装类/副危", "鉴定与 SDS14 一致", "无鉴定或打架=红", True),
    ("EHS", "E3", "易制毒/易制爆/监控化学品/两用物项", "否，或中国出口许可已启动", "疑似未升级=红", True),
    ("EHS", "E4", "抑制剂/稳定剂及报检说明", "不需要，或名称数量可报检", "应加未说明=红", True),
    ("EHS", "E5", "目的国 REACH/CLP/ADR", "本阶段跳过：先出国门", "不适用", False),
    ("实验室", "R1", "粘度/密度/组成（包装选型）", "已提供或有承诺日期", "选包装前缺失=黄", True),
    ("实验室", "R2", "内容物–包装相容性", "仅海关要求时做；未要求=不适用", "海关已要求且无计划=红", True),
    ("实验室", "R3", "出口 13 位码", "与内贸码区分", "未申请阻塞灌装=黄", True),
    ("计划 MP", "M1", "是否该国首出口", "是/否已标记。缓冲天数待整厂流程图", "未判定=红；已标记待天数=黄", True),
    ("计划 MP", "M2", "Filling Order 出口桶码 ≠ 内贸", "系统码已分开", "混码=红", True),
    ("计划 MP", "M3", "排产 vs 包装到货 vs 灌装日 vs 使用鉴定有效期", "灌装日落在证有效期内", "对不上=黄", True),
    ("计划 MP", "M4", "固化剂/稀释剂同期可供、分票规则", "分票身份各自绿", "主剂绿配套红=黄", True),
    ("生产 & QC", "Q1", "罐装单带 UN、现场按单取桶", "知悉出口桶≠内贸桶", "未知=黄（排产前须绿）", True),
    ("生产 & QC", "Q2", "SAP 生产日期精确到日", "QC 确认可做到", "只能到月=红", True),
    ("生产 & QC", "Q3", "小口年切换旧包装报废可执行", "与采购通知一致", "未对齐=黄", True),
    ("仓库", "W1", "UN 包装收货核对 UN+年份、与性能单一致", "有核对标准，异常退采购", "无核对习惯=黄", True),
    ("仓库", "W2", "年切换大口新盖隔离（新盖→3号仓）", "仓知隔离位置", "年切换票未知=黄", True),
    ("仓库", "W3", "出口货与内贸货可分", "不混放、不混贴", "无区分=黄", True),
    ("仓库", "W4", "出库谁贴运输标 vs GHS 标", "与 Order Clearing 一致", "谁贴未定=黄", True),
    ("进出口", "I1", "危化品产地检验（129号）路径是否清楚", "知报检、符合性声明、电子底账", "危化品未定路径=红", True),
    ("进出口", "I2", "使用鉴定申请窗口与齐套资料", "性能单+鉴定+SDS 对独山抬头", "不知如何办证=黄", True),
    ("进出口", "I3", "海运海事 vs 空运 DGD 差异已知", "不把海运证当空运用", "模式不清=红", True),
    ("财务", "F1", "发票抬头、币种（报关用）", "与拟报关路径一致", "未定=黄（报关前红）", True),
    ("财务", "F2", "运费/保险谁付（相对 Incoterms）", "与 B2 一致", "与业务不一致=红", True),
    ("财务", "F3", "目的国进口商 / EORI / 境外 VAT", "本阶段跳过：先出国门", "不适用", False),
]

DEPT_FILL = {
    "业务 / PC": "D6EAF8",
    "采购": "D5F5E3",
    "Order Clearing": "FCF3CF",
    "EHS": "FADBD8",
    "实验室": "E8DAEF",
    "计划 MP": "D6EAF8",
    "生产 & QC": "FDEBD0",
    "仓库": "D5D8DC",
    "进出口": "AED6F1",
    "财务": "F5CBA7",
}

THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Microsoft YaHei", size=10)
BODY_FONT = Font(name="Microsoft YaHei", size=10)
TITLE_FONT = Font(name="Microsoft YaHei", size=16, bold=True, color="1F4E79")
HINT_FONT = Font(name="Microsoft YaHei", size=9, italic=True, color="666666")


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN


def add_yes_no_dv(ws, formula_range):
    dv = DataValidation(type="list", formula1='"☐ 未完成,☑ 已完成,— 不适用"', allow_blank=True)
    dv.error = "请从下拉选择"
    dv.errorTitle = "勾选"
    dv.prompt = "勾选完成状态"
    dv.promptTitle = "完成"
    dv.add(formula_range)
    ws.add_data_validation(dv)


def add_light_dv(ws, formula_range):
    dv = DataValidation(type="list", formula1='"未评,绿,黄,红,不适用"', allow_blank=True)
    dv.error = "请从下拉选择灯"
    dv.errorTitle = "灯"
    dv.add(formula_range)
    ws.add_data_validation(dv)


def light_fills(ws, start_row, end_row, col):
    rng = f"{get_column_letter(col)}{start_row}:{get_column_letter(col)}{end_row}"
    greens = PatternFill("solid", fgColor="C6EFCE")
    yellows = PatternFill("solid", fgColor="FFEB9C")
    reds = PatternFill("solid", fgColor="FFC7CE")
    greys = PatternFill("solid", fgColor="D9D9D9")
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"绿"'], fill=greens))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"黄"'], fill=yellows))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"红"'], fill=reds))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"不适用"'], fill=greys))


def build():
    wb = Workbook()

    # ----- 说明 -----
    ws0 = wb.active
    ws0.title = "使用说明"
    ws0["A1"] = "出口 Initiate 跨部门 Checklist v1"
    ws0["A1"].font = TITLE_FONT
    ws0.merge_cells("A1:F1")
    notes = [
        "Owner：出口部门（物流部下属）。进出口、仓库同属物流部；仓库只答「仓库」行；进出口答 I 段。",
        "触发：业务/PC 已 initiate 出口、尚未锁定 Filling Order。红灯不排产、不订舱。",
        "目标：先合法出中国国门。目的国 REACH/CLP/ADR、境外进口商本阶段勾「— 不适用」。",
        "抬头：一律独山。禁止金山抬头出口，不必再问。",
        "相容性（R2）：仅海关要求时做；未要求选「— 不适用」。",
        "时限与首出口缓冲天数：整厂出口流程图做出后再写入。本表不设回复截止日期。",
        "会签：不强制。出口部门在「票头」勾选「可进入 Filling Order」即放行。会签是否要做另议。",
        "用法：先填「票头」→ 按部门筛选「检查清单」→「完成」列下拉勾选 →「灯」列选绿/黄/红/不适用。",
        "只收书面答复（邮件/S1/S3/PDF）。黄灯尽量填责任人与日期。",
        "对应知识库：knowledge/02-products/export-initiate-checklist.md",
    ]
    for i, t in enumerate(notes, start=3):
        ws0[f"A{i}"] = t
        ws0[f"A{i}"].font = BODY_FONT
        ws0[f"A{i}"].alignment = WRAP
        ws0.merge_cells(f"A{i}:F{i}")
        ws0.row_dimensions[i].height = 22
    ws0.column_dimensions["A"].width = 28
    for col in "BCDEF":
        ws0.column_dimensions[col].width = 18
    ws0.row_dimensions[1].height = 28
    ws0.print_title_rows = "1:1"
    ws0.page_setup.orientation = "landscape"
    ws0.page_setup.fitToPage = True
    ws0.page_setup.fitToWidth = 1
    ws0.page_setup.fitToHeight = 1
    ws0.sheet_properties.tabColor = "1F4E79"

    # ----- 票头 -----
    ws1 = wb.create_sheet("票头")
    ws1.sheet_properties.tabColor = "2E86C1"
    ws1["A1"] = "本票信息（出口部门先填，再发给各部门）"
    ws1["A1"].font = TITLE_FONT
    ws1.merge_cells("A1:D1")
    fields = [
        ("需求号 / SO", ""),
        ("10 位产品码", ""),
        ("目的国", ""),
        ("收货人类型（客户/汉堡/关联方）", ""),
        ("运输方式（海/空/海+空）", ""),
        ("样品或商业", ""),
        ("期望出运周", ""),
        ("出运抬头（锁定）", "独山（不可改为金山）"),
        ("出口部门经办", ""),
        ("发起日期", ""),
        ("是否该国首出口（出口部门判定）", ""),
        ("本票可进入 Filling Order", "☐ 未完成"),
    ]
    ws1["A3"] = "字段"
    ws1["B3"] = "填写"
    ws1["C3"] = "说明"
    style_header(ws1, 3, 3)
    for i, (label, default) in enumerate(fields, start=4):
        ws1.cell(i, 1, label).font = BODY_FONT
        ws1.cell(i, 1).border = THIN
        ws1.cell(i, 1).fill = PatternFill("solid", fgColor="DEEAF6")
        ws1.cell(i, 2, default).font = BODY_FONT
        ws1.cell(i, 2).border = THIN
        ws1.cell(i, 2).alignment = WRAP
    ws1["C4"] = "没有需求号可用 PC 邮件编号"
    ws1["C11"] = "锁定口径，勿改"
    ws1["C14"] = "是 / 否 / 待判定。缓冲天数待流程图"
    ws1["C15"] = "仅出口部门勾选。不要求业务/采购/计划会签"
    for r in range(4, 16):
        ws1.cell(r, 3).font = HINT_FONT
        ws1.cell(r, 3).border = THIN
        for c in range(1, 4):
            ws1.cell(r, c).alignment = WRAP
        ws1.row_dimensions[r].height = 22
    add_yes_no_dv(ws1, "B15")
    first_export_dv = DataValidation(type="list", formula1='"待判定,是,否"', allow_blank=True)
    first_export_dv.add("B14")
    ws1.add_data_validation(first_export_dv)
    mode_dv = DataValidation(type="list", formula1='"海运,空运,海+空,待定"', allow_blank=True)
    mode_dv.add("B8")
    ws1.add_data_validation(mode_dv)
    sample_dv = DataValidation(type="list", formula1='"样品,商业出货"', allow_blank=True)
    sample_dv.add("B9")
    ws1.add_data_validation(sample_dv)
    ws1.column_dimensions["A"].width = 38
    ws1.column_dimensions["B"].width = 36
    ws1.column_dimensions["C"].width = 48
    ws1["A17"] = "红灯汇总（可手工列出编号）"
    ws1["A17"].font = Font(name="Microsoft YaHei", bold=True, size=11)
    ws1.merge_cells("A18:C20")
    ws1["A18"].alignment = Alignment(wrap_text=True, vertical="top")
    ws1["A18"].border = THIN
    ws1.row_dimensions[18].height = 48

    # ----- 清单 -----
    ws = wb.create_sheet("检查清单")
    ws.sheet_properties.tabColor = "117A65"
    headers = [
        "部门",
        "编号",
        "检查项",
        "合格标准",
        "灯规则",
        "完成（勾选）",
        "灯",
        "书面回复 / 附件",
        "责任人",
        "日期",
        "备注",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c, h)
    style_header(ws, 1, len(headers))
    ws.auto_filter.ref = f"A1:K{len(ROWS) + 1}"
    ws.freeze_panes = "C2"

    for i, (dept, code, item, ok, rule, applicable) in enumerate(ROWS, start=2):
        ws.cell(i, 1, dept)
        ws.cell(i, 2, code)
        ws.cell(i, 3, item)
        ws.cell(i, 4, ok)
        ws.cell(i, 5, rule)
        if applicable:
            ws.cell(i, 6, "☐ 未完成")
            ws.cell(i, 7, "未评")
        else:
            ws.cell(i, 6, "— 不适用")
            ws.cell(i, 7, "不适用")
        fill = PatternFill("solid", fgColor=DEPT_FILL.get(dept, "FFFFFF"))
        for c in range(1, 12):
            cell = ws.cell(i, c)
            cell.font = BODY_FONT
            cell.alignment = WRAP if c in (3, 4, 5, 8, 11) else CENTER
            cell.border = THIN
            if c <= 5:
                cell.fill = fill
        ws.row_dimensions[i].height = 36

    last = 1 + len(ROWS)
    add_yes_no_dv(ws, f"F2:F{last}")
    add_light_dv(ws, f"G2:G{last}")
    light_fills(ws, 2, last, 7)

    widths = [16, 8, 42, 40, 32, 16, 12, 28, 12, 12, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:K{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.fitToHeight = 0
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = "1:1"
    ws.oddHeader.left.text = "出口 Initiate Checklist"
    ws.oddHeader.right.text = "独山抬头｜先出国门"

    # ----- 统计 -----
    ws2 = wb.create_sheet("统计")
    ws2.sheet_properties.tabColor = "B7950B"
    ws2["A1"] = "本票灯统计（自动）"
    ws2["A1"].font = TITLE_FONT
    ws2.merge_cells("A1:B1")
    ws2["A3"] = "指标"
    ws2["B3"] = "值"
    style_header(ws2, 3, 2)
    formulas = [
        ("总条数", f"=COUNTA(检查清单!B2:B{last})"),
        ("已完成（☑）", f'=COUNTIF(检查清单!F2:F{last},"☑ 已完成")'),
        ("未完成", f'=COUNTIF(检查清单!F2:F{last},"☐ 未完成")'),
        ("灯=不适用", f'=COUNTIF(检查清单!G2:G{last},"不适用")'),
        ("绿灯", f'=COUNTIF(检查清单!G2:G{last},"绿")'),
        ("黄灯", f'=COUNTIF(检查清单!G2:G{last},"黄")'),
        ("红灯", f'=COUNTIF(检查清单!G2:G{last},"红")'),
        ("未评", f'=COUNTIF(检查清单!G2:G{last},"未评")'),
        ("是否可进 Filling Order", f'=IF(B10>0,"否：有红灯，停排产","出口部门再在票头勾选放行")'),
    ]
    for i, (label, formula) in enumerate(formulas, start=4):
        ws2.cell(i, 1, label).font = BODY_FONT
        ws2.cell(i, 1).border = THIN
        cell = ws2.cell(i, 2, formula)
        cell.font = BODY_FONT
        cell.border = THIN
        cell.alignment = WRAP
        ws2.row_dimensions[i].height = 22
    ws2.column_dimensions["A"].width = 42
    ws2.column_dimensions["B"].width = 42
    ws2["B10"].fill = PatternFill("solid", fgColor="FFC7CE")
    ws2["A15"] = "说明：统计用 COUNTIF，打开 Excel 后即可计算。有任意红灯则不可进 Filling Order；无红灯仍须出口部门在票头勾放行。"
    ws2["A15"].font = HINT_FONT
    ws2.merge_cells("A15:B16")
    ws2["A15"].alignment = WRAP

    wb.properties.title = "出口 Initiate 跨部门 Checklist"
    wb.properties.creator = "Mankiewicz China Logistics"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT} rows={len(ROWS)}")


if __name__ == "__main__":
    build()
