#!/usr/bin/env python3
"""Merge 2022–2025 screenshot extracts with the typed 2026 HAM–SHA snapshot; YoY compare."""
from __future__ import annotations

import csv
from collections import defaultdict
from datetime import date, datetime
from math import ceil
from pathlib import Path
from statistics import mean, median

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path("/workspace")
ARCH = ROOT / "knowledge/06-performance/intl-transport"
SNAP = ARCH / "snapshots"
TYPED = SNAP / "2026-09-07_ham-sha_lvs_bl.csv"
OCR_CSV = ROOT / "deliverables/source/ham_sha_hist_ocr.csv"
SNAPSHOT_ID = "2026-09-09_ham-sha_merged"
AS_OF = date(2026, 9, 7)
XLSX = ROOT / "deliverables/国际运输数据档案_HAM-SHA航行时间.xlsx"

NAVY = "0E2744"
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
BODY = Font(name="Calibri", size=10)
THIN = Border(
    left=Side(style="thin", color="D5DEE6"),
    right=Side(style="thin", color="D5DEE6"),
    top=Side(style="thin", color="D5DEE6"),
    bottom=Side(style="thin", color="D5DEE6"),
)


def parse_d(s: str) -> date:
    s = s.strip().replace("/", "-")
    parts = s.split("-")
    return date(int(parts[0]), int(parts[1]), int(parts[2]))


def norm_bl(s: str) -> str:
    s = str(s).strip().split("/")[0].replace(" ", "")
    return s


def norm_route(s: str) -> str:
    if not s:
        return ""
    r = s.strip().upper().replace(" ", "")
    r = r.replace("TANMED2", "TAN-MED2").replace("TANMED", "TAN-MED")
    r = r.replace("MYTPP", "TPP")
    r = r.replace("PTP", "TPP")
    r = r.replace("1/12SINGAPORE", "HAM-SIN-SHA")
    r = r.replace("1/14SINGAPORE", "HAM-SIN-SHA")
    r = r.replace("1/21SINGAPORE", "HAM-SIN-SHA")
    r = r.replace("1/8SINGAPORE", "HAM-SIN-SHA")
    if "SINGAPORE" in r:
        return "HAM-SIN-SHA"
    return r


def is_direct(route: str) -> str:
    r = (route or "").upper()
    if r in ("", "UNKNOWN"):
        return "U"
    return "Y" if r == "HAM-SHA" else "N"


def load_ocr():
    rows = []
    with OCR_CSV.open() as f:
        for r in csv.DictReader(f):
            bl = norm_bl(r["bl_no"])
            if not bl or not r.get("atd") or not r.get("schedule"):
                continue
            try:
                atd = parse_d(r["atd"])
                sch = parse_d(r["schedule"])
            except Exception:
                continue
            if sch < atd:
                continue
            rows.append(
                {
                    "bl_no": bl,
                    "route": norm_route(r.get("route") or ""),
                    "atd": atd,
                    "schedule": sch,
                    "pri": int(r.get("pri") or 40),
                    "source": r.get("sheet") or "screenshot",
                }
            )
    return rows


def load_typed():
    rows = []
    with TYPED.open() as f:
        for r in csv.DictReader(f):
            atd = parse_d(r["atd"])
            sch = parse_d(r["schedule_date"])
            rows.append(
                {
                    "bl_no": norm_bl(r["bl_no"]),
                    "route": "",
                    "atd": atd,
                    "schedule": sch,
                    "pri": 100,
                    "source": "typed_2026-09-07",
                }
            )
    return rows


def merge(ocr, typed):
    by = {}
    for r in ocr + typed:
        k = r["bl_no"]
        if k not in by:
            by[k] = r
            continue
        cur = by[k]
        if r["pri"] > cur["pri"]:
            route = cur["route"] or r["route"]
            by[k] = {**r, "route": r["route"] or route}
        else:
            if not cur["route"] and r["route"]:
                cur["route"] = r["route"]
            if r["pri"] == cur["pri"] and r["schedule"] > cur["schedule"]:
                by[k] = {**r, "route": r["route"] or cur["route"]}
    out = []
    for r in by.values():
        days = (r["schedule"] - r["atd"]).days
        status = "eta" if r["schedule"] > AS_OF else "schedule_passed"
        out.append(
            {
                "snapshot_id": SNAPSHOT_ID,
                "as_of_date": AS_OF.isoformat(),
                "lane_id": "HAM-SHA",
                "origin": "HAM",
                "destination": "SHA",
                "direction": "import",
                "mode": "ocean",
                "incoterms": "CIF",
                "metric": "latest_vessel_schedule_minus_atd",
                "bl_no": r["bl_no"],
                "route": r["route"] or "",
                "direct": is_direct(r["route"]),
                "atd": r["atd"].isoformat(),
                "atd_year": r["atd"].year,
                "atd_month": r["atd"].strftime("%Y-%m"),
                "schedule_date": r["schedule"].isoformat(),
                "days": days,
                "status": status,
                "source": r["source"],
            }
        )
    out.sort(key=lambda x: (x["atd"], x["bl_no"]))
    return out


def voyages(bl):
    g = defaultdict(list)
    for r in bl:
        g[(r["atd"], r["schedule_date"], r["route"])].append(r)
    rows = []
    for (atd, sch, route), items in sorted(g.items()):
        days = items[0]["days"]
        rows.append(
            {
                "snapshot_id": SNAPSHOT_ID,
                "lane_id": "HAM-SHA",
                "voyage_key": f"{atd}|{sch}|{route}",
                "atd": atd,
                "atd_year": items[0]["atd_year"],
                "atd_month": items[0]["atd_month"],
                "schedule_date": sch,
                "route": route,
                "direct": items[0]["direct"],
                "days": days,
                "bl_count": len(items),
                "status": items[0]["status"],
            }
        )
    return rows


def stats(ds):
    ds = list(ds)
    if not ds:
        return dict(n=0, min="", max="", mean="", median="")
    return dict(
        n=len(ds),
        min=min(ds),
        max=max(ds),
        mean=round(mean(ds), 1),
        median=median(ds),
    )


def coverage(ds, rate=0.8):
    s = sorted(ds)
    if not s:
        return ""
    k = ceil(len(s) * rate)
    return s[k - 1]


def yearly(voy, tickets):
    rows = []
    years = sorted({v["atd_year"] for v in voy})
    for y in years:
        vs = [v for v in voy if v["atd_year"] == y and v["status"] != "eta"]
        ts = [t for t in tickets if t["atd_year"] == y and t["status"] != "eta"]
        vd = [v["days"] for v in vs]
        td = [t["days"] for t in ts]
        dv = [v["days"] for v in vs if v["direct"] == "Y"]
        nv = [v["days"] for v in vs if v["direct"] == "N"]
        sv = stats(vd)
        st = stats(td)
        sd = stats(dv)
        rows.append(
            {
                "atd_year": y,
                "voyage_n": sv["n"],
                "voyage_min": sv["min"],
                "voyage_max": sv["max"],
                "voyage_mean": sv["mean"],
                "voyage_median": sv["median"],
                "voyage_p80": coverage(vd, 0.8),
                "direct_voyage_n": sd["n"],
                "direct_voyage_mean": sd["mean"],
                "transship_voyage_n": stats(nv)["n"],
                "transship_voyage_mean": stats(nv)["mean"],
                "ticket_n": st["n"],
                "ticket_min": st["min"],
                "ticket_max": st["max"],
                "ticket_mean": st["mean"],
                "ticket_median": st["median"],
            }
        )
    return rows


def monthly(voy):
    by = defaultdict(list)
    for v in voy:
        if v["status"] == "eta":
            continue
        by[v["atd_month"]].append(v["days"])
    rows = []
    for m in sorted(by):
        s = stats(by[m])
        rows.append(
            {
                "atd_month": m,
                "atd_year": int(m[:4]),
                "n": s["n"],
                "min": s["min"],
                "max": s["max"],
                "mean": s["mean"],
                "median": s["median"],
            }
        )
    return rows


def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def sheet(wb, name, rows, widths):
    ws = wb.create_sheet(name)
    fields = list(rows[0].keys())
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{len(rows)+1}"
    for c, h in enumerate(fields, 1):
        cell = ws.cell(1, c, h)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center")
    for i, r in enumerate(rows, 2):
        for c, h in enumerate(fields, 1):
            cell = ws.cell(i, c, r[h])
            cell.font = BODY
            cell.border = THIN
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    return ws


def write_xlsx(bl, voy, year, month):
    wb = Workbook()
    intro = wb.active
    intro.title = "使用说明"
    intro["A1"] = "国际运输数据档案 — HAM→SHA 合并原始表 + 年度比较"
    intro["A1"].font = Font(bold=True, size=16, color=NAVY, name="Calibri")
    intro.merge_cells("A1:B1")
    notes = [
        ("快照", SNAPSHOT_ID),
        ("口径", "days = Latest Vessel Schedule − ATD（日历天）；年 = ATD 年"),
        ("计划单元", "航次 = 同一 ATD + 船期日 + 路由，一航次一权"),
        ("去重", "主键 B/L No.；与 2026 打字底表冲突时以 2026 底表为准，路由可从截图补"),
        ("2022–2025 来源", "用户附件截图转录；个别票号/日期可能有识读误差"),
        ("2026", "typed extract 107 票（含 2025-11/12 开航）优先"),
        ("ETA", f"船期晚于 {AS_OF.isoformat()} 的 2026-07/08 不进年度对比"),
        ("票数", f"{len(bl)} 去重后；航次 {len(voy)}"),
        ("知识档案", "knowledge/06-performance/intl-transport/"),
    ]
    intro["A3"] = "项"
    intro["B3"] = "内容"
    intro["A3"].fill = HEAD_FILL
    intro["B3"].fill = HEAD_FILL
    intro["A3"].font = HEAD_FONT
    intro["B3"].font = HEAD_FONT
    for i, (k, v) in enumerate(notes, 4):
        intro.cell(i, 1, k).font = Font(bold=True, name="Calibri", size=10)
        intro.cell(i, 2, v).font = BODY
    intro.column_dimensions["A"].width = 18
    intro.column_dimensions["B"].width = 92

    # drop snapshot-only cols for excel readability
    bl_x = [{k: r[k] for k in ("bl_no", "route", "direct", "atd", "atd_year", "atd_month", "schedule_date", "days", "status", "source")} for r in bl]
    voy_x = [{k: r[k] for k in ("atd", "atd_year", "atd_month", "schedule_date", "route", "direct", "days", "bl_count", "status")} for r in voy]
    sheet(wb, "原始数据_票", bl_x, [14, 22, 8, 12, 10, 10, 14, 8, 16, 22])
    sheet(wb, "原始数据_航次", voy_x, [12, 10, 10, 14, 22, 8, 8, 10, 16])
    sheet(wb, "年度比较", year, [10, 12, 12, 12, 14, 14, 12, 16, 18, 12, 12, 12, 12, 14])
    sheet(wb, "分月航次", month, [12, 10, 8, 8, 8, 10, 10])
    # remove default if extra
    wb.save(XLSX)


def update_catalog(bl, voy):
    cat_path = ARCH / "catalog.csv"
    rows = list(csv.DictReader(cat_path.open()))
    fields = list(rows[0].keys()) if rows else [
        "snapshot_id", "as_of_date", "lane_id", "metric", "n_bl", "n_voyage",
        "atd_from", "atd_to", "n_eta_voyage", "notes",
    ]
    rows = [r for r in rows if r["snapshot_id"] != SNAPSHOT_ID]
    rows.append(
        {
            "snapshot_id": SNAPSHOT_ID,
            "as_of_date": AS_OF.isoformat(),
            "lane_id": "HAM-SHA",
            "metric": "latest_vessel_schedule_minus_atd",
            "n_bl": len(bl),
            "n_voyage": len(voy),
            "atd_from": min(r["atd"] for r in bl),
            "atd_to": max(r["atd"] for r in bl),
            "n_eta_voyage": sum(1 for v in voy if v["status"] == "eta"),
            "notes": "Merged screenshot 2022–2025 + typed 2026 extract; dedupe by B/L.",
        }
    )
    with cat_path.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def main():
    ocr = load_ocr()
    typed = load_typed()
    bl = merge(ocr, typed)
    voy = voyages(bl)
    year = yearly(voy, bl)
    month = monthly(voy)
    write_csv(SNAP / f"{SNAPSHOT_ID}_bl.csv", bl)
    write_csv(SNAP / f"{SNAPSHOT_ID}_voyage.csv", voy)
    write_csv(SNAP / f"{SNAPSHOT_ID}_yearly.csv", year)
    write_csv(SNAP / f"{SNAPSHOT_ID}_monthly.csv", month)
    write_xlsx(bl, voy, year, month)
    update_catalog(bl, voy)
    print("tickets", len(bl), "voyages", len(voy))
    print("typed overlays", sum(1 for r in bl if r["source"] == "typed_2026-09-07"))
    print("YEARLY")
    for r in year:
        print(r)


if __name__ == "__main__":
    main()
