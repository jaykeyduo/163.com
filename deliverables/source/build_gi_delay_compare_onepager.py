#!/usr/bin/env python3
"""A4 landscape: two warehouse-ETA delay clocks vs German GI date."""
from __future__ import annotations

import csv
from collections import defaultdict
from datetime import date, datetime, timedelta
from math import ceil
from pathlib import Path
from statistics import mean, median

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import load_workbook
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

pdfmetrics.registerFont(
    TTFont("CN", "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf")
)

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "compare_gi_warehouse_delay.xlsx"
DAILY_CSV = ROOT / "compare_gi_warehouse_delay_daily.csv"
CHART_PNG = Path("/tmp/gi_delay_two_clocks.png")
OUT = Path("/workspace/deliverables/GI_warehouse_delay_two_clocks.pdf")

NAVY = HexColor("#0E2744")
RED = HexColor("#C0392B")
GREEN = HexColor("#1F6B4A")
MUTED = HexColor("#5B6B7C")
INK = HexColor("#24384A")
LINE = HexColor("#D5DEE6")
BG = HexColor("#F7F9FB")
BG_BAD = HexColor("#FDECEA")
BG_OK = HexColor("#EAF6EF")
BG_NOTE = HexColor("#EEF3F8")

PL_COLOR = "#1B4F8A"
AGI_COLOR = "#C0392B"

FONT_PATH = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"
font_manager.fontManager.addfont(FONT_PATH)
plt.rcParams["font.family"] = "WenQuanYi Micro Hei"
plt.rcParams["axes.unicode_minus"] = False


def xldate(n):
    if isinstance(n, datetime):
        return n.date()
    if isinstance(n, date):
        return n
    if isinstance(n, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=int(n))).date()
    return None


def parse_delay(v):
    if isinstance(v, (int, float)):
        return float(v)
    if v is None:
        return None
    s = str(v).strip()
    if s in ("", "(空白)", "空白", "-", "总计"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def coverage(vals, rate=0.80):
    s = sorted(vals)
    k = ceil(len(s) * rate)
    return s[k - 1]


def load_rows():
    wb = load_workbook(XLSX, data_only=True)
    ws = wb.active
    rows = []
    dropped = {"total": 0, "pl_outlier": 0, "no_date": 0}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        d = xldate(r[0])
        if d is None:
            dropped["no_date"] += 1
            continue
        pl = parse_delay(r[3])
        agi = parse_delay(r[4])
        if pl is not None and pl < -100:
            dropped["pl_outlier"] += 1
            pl = None
        rows.append(
            {
                "date": d,
                "sd": str(r[1]) if r[1] is not None else "",
                "dlv": str(r[2]) if r[2] is not None else "",
                "pl": pl,
                "agi": agi,
            }
        )
    return rows, dropped


def delivery_grain(rows):
    by = defaultdict(list)
    for x in rows:
        by[x["dlv"]].append(x)
    out = []
    for dlv, items in by.items():
        pls = [x["pl"] for x in items if x["pl"] is not None]
        ags = [x["agi"] for x in items if x["agi"] is not None]
        out.append(
            {
                "dlv": dlv,
                "date": min(x["date"] for x in items),
                "n_so": len(items),
                "pl": mean(pls) if pls else None,
                "agi": mean(ags) if ags else None,
            }
        )
    return out


def daily_series(dlvs):
    by = defaultdict(lambda: {"pl": [], "agi": [], "n_dlv": 0, "n_so": 0})
    for x in dlvs:
        k = x["date"]
        by[k]["n_dlv"] += 1
        by[k]["n_so"] += x["n_so"]
        if x["pl"] is not None:
            by[k]["pl"].append(x["pl"])
        if x["agi"] is not None:
            by[k]["agi"].append(x["agi"])
    series = []
    for d in sorted(by):
        pl = by[d]["pl"]
        ag = by[d]["agi"]
        if not pl or not ag:
            continue
        series.append(
            {
                "date": d.isoformat(),
                "n_dlv": by[d]["n_dlv"],
                "n_so": by[d]["n_so"],
                "pl_mean": mean(pl),
                "agi_mean": mean(ag),
                "gap": mean(pl) - mean(ag),
            }
        )
    return series


def monthly(dlvs, year=2026):
    mon = defaultdict(lambda: {"pl": [], "agi": []})
    for x in dlvs:
        if x["date"].year != year:
            continue
        key = x["date"].month
        if x["pl"] is not None:
            mon[key]["pl"].append(x["pl"])
        if x["agi"] is not None:
            mon[key]["agi"].append(x["agi"])
    out = {}
    for m, v in mon.items():
        out[m] = {
            "n": len(v["agi"]),
            "pl": mean(v["pl"]) if v["pl"] else None,
            "agi": mean(v["agi"]) if v["agi"] else None,
        }
    return out


def wrap(text: str, width: int) -> list[str]:
    words = text.split()
    lines, cur = [], ""
    for w in words:
        trial = (cur + " " + w).strip()
        if len(trial) <= width:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def draw_chart(series):
    xs = [datetime.fromisoformat(r["date"]).date() for r in series]
    pl = [r["pl_mean"] for r in series]
    ag = [r["agi_mean"] for r in series]
    ns = [r["n_dlv"] for r in series]

    fig, ax = plt.subplots(figsize=(10.6, 3.55), dpi=180)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    ax.fill_between(xs, pl, ag, color="#C0392B", alpha=0.10, linewidth=0)
    ax.plot(xs, pl, color=PL_COLOR, lw=2.0, marker="o", ms=3.6, label="相对 Pl.arr.date（SAP 计划到仓）")
    ax.plot(xs, ag, color=AGI_COLOR, lw=2.0, marker="s", ms=3.4, label="相对 AGI+72（发货日+72天）")
    ax.axhline(0, color="#1F6B4A", lw=0.9, ls="--", alpha=0.85)
    ax.text(xs[0], 1.2, "0 = 准时", color="#1F6B4A", fontsize=8)
    # highlight heavy GI days
    for x, ypl, yag, n in zip(xs, pl, ag, ns):
        if n >= 25:
            ax.annotate(
                f"n={n}",
                xy=(x, max(ypl, yag)),
                xytext=(0, 8),
                textcoords="offset points",
                ha="center",
                fontsize=7,
                color="#5B6B7C",
            )
    ax.set_ylabel("延迟天数（>0 晚到）")
    ax.set_xlabel("Actual goods movement date（德国物流发货日）")
    ax.set_ylim(-4, 42)
    ax.yaxis.set_major_locator(plt.MultipleLocator(5))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%m/%d"))
    ax.xaxis.set_major_locator(mdates.WeekdayLocator(byweekday=mdates.MO))
    ax.grid(True, axis="y", color="#E6EBF0", lw=0.7)
    ax.grid(True, axis="x", color="#F0F3F6", lw=0.5)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color("#D5DEE6")
    ax.spines["bottom"].set_color("#D5DEE6")
    ax.legend(frameon=False, loc="upper right", fontsize=8.5)
    fig.tight_layout(pad=0.35)
    fig.savefig(CHART_PNG, dpi=180)
    plt.close(fig)


def card(c, x, y, w, h, fill, title, value, sub, value_color):
    c.setFillColor(fill)
    c.roundRect(x, y, w, h, 2.2 * mm, fill=1, stroke=0)
    c.setFillColor(MUTED)
    c.setFont("Helvetica", 6.4)
    c.drawString(x + 2.4 * mm, y + h - 6.0 * mm, title)
    c.setFillColor(value_color)
    c.setFont("Helvetica-Bold", 15)
    c.drawString(x + 2.4 * mm, y + 8.4 * mm, value)
    c.setFillColor(INK)
    c.setFont("Helvetica", 6.1)
    c.drawString(x + 2.4 * mm, y + 3.0 * mm, sub)


def main():
    rows, dropped = load_rows()
    dlvs = delivery_grain(rows)
    d2026 = [x for x in dlvs if x["date"].year == 2026]
    series = daily_series(d2026)
    with DAILY_CSV.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["date", "n_dlv", "n_so", "pl_mean", "agi_mean", "gap"])
        w.writeheader()
        for r in series:
            w.writerow(
                {
                    **r,
                    "pl_mean": round(r["pl_mean"], 2),
                    "agi_mean": round(r["agi_mean"], 2),
                    "gap": round(r["gap"], 2),
                }
            )

    pl = [x["pl"] for x in d2026 if x["pl"] is not None]
    ag = [x["agi"] for x in d2026 if x["agi"] is not None]
    both = [x for x in d2026 if x["pl"] is not None and x["agi"] is not None]
    gap = [x["pl"] - x["agi"] for x in both]
    mon = monthly(d2026)
    n_so = sum(x["n_so"] for x in d2026)
    n_early_pl = sum(1 for v in pl if v < 0)
    share_pl_worse = 100 * sum(1 for x in both if x["pl"] > x["agi"]) / len(both)

    draw_chart(series)

    page_w, page_h = landscape(A4)
    c = canvas.Canvas(str(OUT), pagesize=landscape(A4))
    m = 9 * mm

    c.setFillColor(NAVY)
    c.rect(0, page_h - 22 * mm, page_w, 22 * mm, fill=1, stroke=0)
    c.setFillColor(RED)
    c.rect(0, page_h - 22 * mm, 3.2 * mm, 22 * mm, fill=1, stroke=0)
    c.setFillColor(white)
    c.setFont("Helvetica-Bold", 14.2)
    c.drawString(m + 2 * mm, page_h - 8.4 * mm, "Two clocks, one GI date — warehouse ETA delay")
    c.setFillColor(white)
    c.setFont("Helvetica", 8)
    c.drawString(
        m + 2 * mm,
        page_h - 13.6 * mm,
        "German GI date  |  warehouse ETA delay  |  Pl.arr.date vs AGI+72  |  delivery-equal  |  Jun-Aug 2026",
    )
    c.setFont("Helvetica", 7.4)
    c.drawString(m + 2 * mm, page_h - 18.6 * mm, "X = Actual goods movement date    Y = delay days    >0 late, <0 early")
    c.setFont("Helvetica", 7.5)
    c.drawRightString(page_w - m, page_h - 9.0 * mm, "one page  ·  logistics")
    c.drawRightString(page_w - m, page_h - 15.4 * mm, f"{len(d2026)} deliveries  ·  {n_so} SO")

    kpi_y = page_h - 22 * mm - 27.5 * mm
    kpi_h = 25.5 * mm
    gap_w = 3 * mm
    usable = page_w - 2 * m
    cw = (usable - 4 * gap_w) / 5

    card(
        c,
        m,
        kpi_y,
        cw,
        kpi_h,
        BG,
        "vs Pl.arr.date  mean",
        f"{mean(pl):.1f} d",
        f"med {median(pl):.0f}d   80% {coverage(pl)}d   n={len(pl)}",
        NAVY,
    )
    card(
        c,
        m + cw + gap_w,
        kpi_y,
        cw,
        kpi_h,
        BG,
        "vs AGI+72  mean",
        f"{mean(ag):.1f} d",
        f"med {median(ag):.0f}d   80% {coverage(ag)}d   n={len(ag)}",
        RED,
    )
    card(
        c,
        m + 2 * (cw + gap_w),
        kpi_y,
        cw,
        kpi_h,
        BG_BAD,
        "Pl.arr minus AGI+72",
        f"+{mean(gap):.1f} d",
        f"Pl.arr clock is tighter on {share_pl_worse:.0f}% of deliveries",
        RED,
    )
    card(
        c,
        m + 3 * (cw + gap_w),
        kpi_y,
        cw,
        kpi_h,
        BG_OK,
        "AGI+72 early deliveries",
        "0",
        f"min delay {min(ag):.0f}d — 72-day clock never beaten",
        GREEN,
    )
    card(
        c,
        m + 4 * (cw + gap_w),
        kpi_y,
        cw,
        kpi_h,
        BG_NOTE,
        "Jun → Aug  (AGI+72 mean)",
        f"{mon[6]['agi']:.0f}→{mon[8]['agi']:.0f} d",
        f"Jun {mon[6]['agi']:.1f}  Jul {mon[7]['agi']:.1f}  Aug {mon[8]['agi']:.1f}  ·  Aug still ETA",
        NAVY,
    )

    img = ImageReader(str(CHART_PNG))
    chart_y = 52 * mm
    chart_h = kpi_y - 5 * mm - chart_y
    c.drawImage(img, m, chart_y, width=usable, height=chart_h, preserveAspectRatio=True, mask="auto")

    box_y = 8.5 * mm
    box_h = 41.5 * mm
    bw = (usable - 2 * gap_w) / 3
    boxes = [
        (
            BG_NOTE,
            NAVY,
            "HOW TO READ",
            [
                "Delay = expected warehouse arrival − clock. Same ETA, two clocks.",
                "Blue = SAP Pl.arr.date. Red = Actual GI date + 72 calendar days.",
                "One point per GI date = mean of deliveries that left DE that day (not SO-weighted).",
                f"Dropped 1 SO at −272d and 1 Sep-2025 row without Pl.arr. Source n={len(rows)} SO / {len(dlvs)} deliveries.",
            ],
        ),
        (
            BG,
            NAVY,
            "WHAT THE GAP MEANS",
            [
                f"Pl.arr delay averages {mean(gap):.1f}d higher than AGI+72. Planned arrival is ~3 days tighter than the 72-day rule.",
                f"Only {n_early_pl} deliveries beat Pl.arr.date. None beat AGI+72 (floor +{min(ag):.0f}d).",
                f"June GI: Pl.arr {mon[6]['pl']:.1f}d / AGI+72 {mon[6]['agi']:.1f}d. August: {mon[8]['pl']:.1f} / {mon[8]['agi']:.1f}.",
                "The two lines move together. The clock choice changes the level, not the shape.",
            ],
        ),
        (
            BG_BAD,
            RED,
            "DO NOT READ AUGUST AS RECOVERY",
            [
                "August GI warehouse dates are still ETAs. Less delay can mean less time for the ETA to slip.",
                "Heavy GI days (n≥25 labelled) still sit on the same two-clock gap.",
                "This is warehouse-ETA delay, not HAM→SHA sailing days, and not door-to-Dushan.",
                "Factory 55d sailing LT and AGI+72 are different clocks — do not swap them.",
            ],
        ),
    ]
    for i, (fill, accent, title, lines) in enumerate(boxes):
        x = m + i * (bw + gap_w)
        c.setFillColor(fill)
        c.roundRect(x, box_y, bw, box_h, 2 * mm, fill=1, stroke=0)
        c.setFillColor(accent)
        c.rect(x, box_y, 1.6 * mm, box_h, fill=1, stroke=0)
        c.setFont("Helvetica-Bold", 7.1)
        c.drawString(x + 3.8 * mm, box_y + box_h - 6.2 * mm, title)
        c.setFillColor(INK)
        c.setFont("Helvetica", 6.05)
        ty = box_y + box_h - 11.4 * mm
        for line in lines:
            for wline in wrap(line, 52):
                c.drawString(x + 3.8 * mm, ty, wline)
                ty -= 3.25 * mm
            ty -= 0.45 * mm

    c.showPage()
    c.save()
    print("wrote", OUT)
    print("daily", DAILY_CSV)
    print("chart", CHART_PNG)
    print("2026 deliveries", len(d2026), "SO", n_so)
    print("pl mean", round(mean(pl), 2), "agi mean", round(mean(ag), 2), "gap", round(mean(gap), 2))
    print("dropped", dropped)
    print("months", {k: {kk: (round(vv, 2) if isinstance(vv, float) else vv) for kk, vv in v.items()} for k, v in mon.items()})


if __name__ == "__main__":
    main()
