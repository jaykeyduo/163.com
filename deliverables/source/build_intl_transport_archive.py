#!/usr/bin/env python3
"""Build the international-transport archive snapshot from the HAM–SHA B/L extract."""
from __future__ import annotations

import csv
from collections import defaultdict
from datetime import date, datetime
from math import ceil
from pathlib import Path
from statistics import mean, median

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path("/workspace")
SRC_BL = ROOT / "deliverables/source/ham_sha_latest_schedule_bl.csv"
ARCH = ROOT / "knowledge/06-performance/intl-transport"
SNAP_DIR = ARCH / "snapshots"
XLSX = ROOT / "deliverables/国际运输数据档案_HAM-SHA航行时间.xlsx"

SNAPSHOT_ID = "2026-09-07_ham-sha_lvs"
AS_OF = date(2026, 9, 7)
LANE_ID = "HAM-SHA"
METRIC = "latest_vessel_schedule_minus_atd"


def parse_iso(s: str) -> date:
    return datetime.fromisoformat(s).date()


def coverage(days, rate: float):
    s = sorted(days)
    k = ceil(len(s) * rate)
    return k, s[k - 1]


def load_bl():
    rows = []
    with SRC_BL.open() as f:
        for r in csv.DictReader(f):
            atd = parse_iso(r["ATD"])
            sch = parse_iso(r["Latest Vessel Schedule"])
            days = int(r["Days"])
            assert (sch - atd).days == days, r
            rows.append(
                {
                    "snapshot_id": SNAPSHOT_ID,
                    "as_of_date": ASOF_S,
                    "lane_id": LANE_ID,
                    "origin": "HAM",
                    "destination": "SHA",
                    "direction": "import",
                    "mode": "ocean",
                    "incoterms": "CIF",
                    "metric": METRIC,
                    "bl_no": r["B/L No."],
                    "atd": atd.isoformat(),
                    "atd_month": atd.strftime("%Y-%m"),
                    "schedule_date": sch.isoformat(),
                    "days": days,
                    "status": "eta" if sch > AS_OF else "schedule_passed",
                    "source": "user_extract_2026-09-07_v2",
                }
            )
    return rows


ASOF_S = AS_OF.isoformat()


def derive_voyages(bl_rows):
    groups = defaultdict(list)
    for r in bl_rows:
        groups[(r["atd"], r["schedule_date"])].append(r)
    voyages = []
    for (atd, sch), items in sorted(groups.items()):
        days = items[0]["days"]
        assert all(x["days"] == days for x in items)
        voyages.append(
            {
                "snapshot_id": SNAPSHOT_ID,
                "as_of_date": ASOF_S,
                "lane_id": LANE_ID,
                "metric": METRIC,
                "voyage_key": f"{atd}|{sch}",
                "atd": atd,
                "atd_month": items[0]["atd_month"],
                "schedule_date": sch,
                "days": days,
                "bl_count": len(items),
                "status": items[0]["status"],
                "bl_nos": ";".join(x["bl_no"] for x in items),
            }
        )
    return voyages


def monthly_rows(bl_rows, voyages):
    out = []

    def emit(grain, month, ds, status):
        out.append(
            {
                "snapshot_id": SNAPSHOT_ID,
                "as_of_date": ASOF_S,
                "lane_id": LANE_ID,
                "metric": METRIC,
                "grain": grain,
                "atd_month": month,
                "n": len(ds),
                "min_days": min(ds),
                "max_days": max(ds),
                "mean_days": round(mean(ds), 4),
                "median_days": median(ds),
                "status": status,
            }
        )

    by_m_v = defaultdict(list)
    by_m_t = defaultdict(list)
    st_m = defaultdict(set)
    for v in voyages:
        by_m_v[v["atd_month"]].append(v["days"])
        st_m[v["atd_month"]].add(v["status"])
    for r in bl_rows:
        by_m_t[r["atd_month"]].append(r["days"])
    for m in sorted(by_m_v):
        status = "eta" if st_m[m] == {"eta"} else "mixed" if len(st_m[m]) > 1 else "schedule_passed"
        emit("voyage", m, by_m_v[m], status)
        emit("ticket", m, by_m_t[m], status)
    return out


def coverage_rows(voyages):
    complete = [v["days"] for v in voyages if v["status"] != "eta"]
    windows = {
        "all_incl_eta": [v["days"] for v in voyages],
        "nov_jun_schedule_passed": complete,
        "feb_may": [
            v["days"]
            for v in voyages
            if v["atd"][:7] in ("2026-02", "2026-03", "2026-04", "2026-05")
        ],
        "apr_jun": [
            v["days"]
            for v in voyages
            if v["status"] != "eta" and v["atd"][:7] in ("2026-04", "2026-05", "2026-06")
        ],
    }
    out = []
    for name, ds in windows.items():
        for rate in (0.80, 0.85, 0.90, 0.95):
            k, L = coverage(ds, rate)
            out.append(
                {
                    "snapshot_id": SNAPSHOT_ID,
                    "lane_id": LANE_ID,
                    "metric": METRIC,
                    "grain": "voyage",
                    "window": name,
                    "n": len(ds),
                    "coverage": rate,
                    "k": k,
                    "L_days": L,
                    "mean_days": round(mean(ds), 4),
                    "median_days": median(ds),
                    "min_days": min(ds),
                    "max_days": max(ds),
                }
            )
    return out


def write_csv(path: Path, rows, fieldnames):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)


def write_xlsx(bl, voy, monthly, cov, catalog, lanes):
    wb = Workbook()

    thin = Border(
        left=Side(style="thin", color="D5DEE6"),
        right=Side(style="thin", color="D5DEE6"),
        top=Side(style="thin", color="D5DEE6"),
        bottom=Side(style="thin", color="D5DEE6"),
    )
    head_fill = PatternFill("solid", fgColor="0E2744")
    head_font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    body_font = Font(name="Calibri", size=10)

    def sheet(name, rows, fields, widths):
        ws = wb.create_sheet(name)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{len(rows)+1}"
        for c, h in enumerate(fields, 1):
            cell = ws.cell(1, c, h)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center")
        for i, r in enumerate(rows, 2):
            for c, h in enumerate(fields, 1):
                cell = ws.cell(i, c, r[h])
                cell.font = body_font
                cell.border = thin
                cell.alignment = Alignment(vertical="center")
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        return ws

    intro = wb.active
    intro.title = "使用说明"
    intro["A1"] = "国际运输数据档案 — HAM→SHA 航行时间"
    intro["A1"].font = Font(bold=True, size=16, color="0E2744", name="Calibri")
    intro.merge_cells("A1:B1")
    lines = [
        ("用途", "后续对比、滚动分析的业务底表。不要覆盖历史快照，只追加。"),
        ("线路", "HAM→SHA 进口海运 CIF（汉堡→上海）。门到独山不含在本表。"),
        ("口径", "days = Latest Vessel Schedule − ATD（日历天）。"),
        ("计划单元", "航次 = 同一 ATD + 同一船期日，一航次一权。按票表只作对照。"),
        ("截止日", "2026-09-07。status=eta 表示船期日晚于截止日，对比时默认剔除。"),
        ("快照ID", SNAPSHOT_ID),
        ("票数 / 航次", f"{len(bl)} / {len(voy)}"),
        ("ATD 范围", "2025-11-23 ~ 2026-08-07"),
        ("知识档案", "knowledge/06-performance/intl-transport/"),
        ("计划算法", "knowledge/06-performance/ham-sha-sailing-lead-time.md"),
        ("禁止", "用 MIN / 全样本 AVG / 单票 MAX 当 sailing LT；ETA 月不当恢复。"),
    ]
    intro["A3"] = "项"
    intro["B3"] = "内容"
    intro["A3"].font = head_font
    intro["B3"].font = head_font
    intro["A3"].fill = head_fill
    intro["B3"].fill = head_fill
    for i, (k, v) in enumerate(lines, 4):
        intro.cell(i, 1, k).font = Font(bold=True, name="Calibri", size=10)
        intro.cell(i, 2, v).font = body_font
    intro.column_dimensions["A"].width = 16
    intro.column_dimensions["B"].width = 88

    sheet("线路目录", lanes, list(lanes[0].keys()), [14, 10, 16, 14, 22, 12, 10, 12, 28, 12, 40])
    sheet("快照目录", catalog, list(catalog[0].keys()), [28, 14, 12, 36, 10, 10, 14, 14, 18, 40])
    sheet("HAM-SHA票", bl, list(bl[0].keys()), [26, 12, 12, 8, 12, 10, 8, 10, 32, 14, 12, 12, 14, 8, 16, 28])
    sheet("HAM-SHA航次", voy, list(voy[0].keys()), [26, 12, 12, 32, 22, 12, 12, 14, 8, 10, 16, 40])
    sheet("HAM-SHA月汇总", monthly, list(monthly[0].keys()), [26, 12, 12, 32, 10, 10, 8, 10, 10, 12, 12, 16])
    sheet("覆盖率窗口", cov, list(cov[0].keys()), [26, 12, 32, 10, 24, 8, 10, 8, 10, 12, 12, 10, 10])

    wb.save(XLSX)


def main():
    bl = load_bl()
    voy = derive_voyages(bl)
    monthly = monthly_rows(bl, voy)
    cov = coverage_rows(voy)

    lanes = [
        {
            "lane_id": "HAM-SHA",
            "origin": "HAM",
            "origin_name": "Hamburg",
            "destination": "SHA",
            "dest_name": "Shanghai (Yangshan)",
            "direction": "import",
            "mode": "ocean",
            "incoterms": "CIF",
            "first_snapshot_id": SNAPSHOT_ID,
            "status": "active",
            "notes": "Germany-nominated carrier. Sailing/port-to-port days only.",
        },
        {
            "lane_id": "FRA-PVG",
            "origin": "FRA",
            "origin_name": "Frankfurt",
            "destination": "PVG",
            "dest_name": "Shanghai Pudong",
            "direction": "import",
            "mode": "air",
            "incoterms": "CIF",
            "first_snapshot_id": "",
            "status": "awaiting_extract",
            "notes": "Lane exists in knowledge/03-network/network-overview.md. No lead-time extract yet — do not invent days.",
        },
    ]

    catalog = [
        {
            "snapshot_id": SNAPSHOT_ID,
            "as_of_date": ASOF_S,
            "lane_id": LANE_ID,
            "metric": METRIC,
            "n_bl": len(bl),
            "n_voyage": len(voy),
            "atd_from": min(r["atd"] for r in bl),
            "atd_to": max(r["atd"] for r in bl),
            "n_eta_voyage": sum(1 for v in voy if v["status"] == "eta"),
            "notes": "v2 user extract. Jul–Aug = ETA as of 7 Sep 2026. Supersedes 60-B/L Apr cut.",
        }
    ]

    SNAP_DIR.mkdir(parents=True, exist_ok=True)
    write_csv(ARCH / "lanes.csv", lanes, list(lanes[0].keys()))
    write_csv(ARCH / "catalog.csv", catalog, list(catalog[0].keys()))
    write_csv(SNAP_DIR / f"{SNAPSHOT_ID}_bl.csv", bl, list(bl[0].keys()))
    write_csv(SNAP_DIR / f"{SNAPSHOT_ID}_voyage.csv", voy, list(voy[0].keys()))
    write_csv(SNAP_DIR / f"{SNAPSHOT_ID}_monthly.csv", monthly, list(monthly[0].keys()))
    write_csv(SNAP_DIR / f"{SNAPSHOT_ID}_coverage.csv", cov, list(cov[0].keys()))
    write_xlsx(bl, voy, monthly, cov, catalog, lanes)

    print("bl", len(bl), "voy", len(voy), "monthly", len(monthly), "coverage", len(cov))
    print("xlsx", XLSX)


if __name__ == "__main__":
    main()
